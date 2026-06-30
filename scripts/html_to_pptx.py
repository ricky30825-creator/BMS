#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
교정된 와이어프레임 HTML(AI_배터리_열폭주_관제_시스템_v2.html)을
편집 가능한 PPTX(발표자료/와이어프레임_v2.pptx)로 변환한다.

방식: 헤드리스 Chrome(Playwright)으로 각 화면(1440x900) 안의 요소들의 실제 렌더 좌표/스타일을
추출 → python-pptx 네이티브 도형(사각형/텍스트박스/그림)으로 1:1 재구성.
→ 스크린샷 이미지가 아니라 PowerPoint에서 직접 편집 가능한 객체.
"""
import re, io, os, sys
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "web", "AI_배터리_열폭주_관제_시스템_v2.html")
OUT = os.path.join(ROOT, "발표자료", "와이어프레임_v2.pptx")
RW, RH = 1440.0, 900.0          # 렌더러 픽셀 크기 (16:10)
SLIDE_W_IN = 13.333
EMU_PER_PX = SLIDE_W_IN * 914400 / RW   # 균일 스케일 (가로/세로 동일)
# 슬라이드 캔버스는 108px/in (1440px→13.333in). 폰트 pt = px * 72/108 로 변환해야
# 레이아웃과 글자 크기가 일치(0.75는 96dpi 가정이라 글자가 ~12.5% 커져 박스를 넘쳤음).
FONT_PT_PER_PX = 72.0 / (RW / SLIDE_W_IN)   # = 0.6667
# Chrome 와이어프레임이 -apple-system(=Apple SD Gothic Neo)로 렌더했으므로 동일 폰트 사용 → 폭 일치
KFONT = "Apple SD Gothic Neo"

# 화면별 fit-to-bounds: 콘텐츠 실제 범위를 계산해 슬라이드 안에 여백 두고 맞춤.
# 기본 0.92로 여백을 주되, 콘텐츠가 더 크면(세로로 긴 화면 등) 더 축소해 경계 이탈 방지.
SCALE_BASE = 0.92
MARGIN_PX = 36                  # 슬라이드 내부 여백(1440x900 기준 px)
CUR = {'sc': SCALE_BASE, 'ox': 0.0, 'oy': 0.0}   # 현재 슬라이드의 스케일/오프셋

def fit_screen(boxes):
    """이 화면 콘텐츠가 슬라이드(여백 포함)에 들어가도록 스케일/오프셋 계산."""
    xs0=[b['x'] for b in boxes]; ys0=[b['y'] for b in boxes]
    xs1=[b['x']+b['w'] for b in boxes]; ys1=[b['y']+b['h'] for b in boxes]
    minX=min([0.0]+xs0); minY=min([0.0]+ys0)
    maxX=max([RW]+xs1);  maxY=max([RH]+ys1)
    cW=max(maxX-minX,1.0); cH=max(maxY-minY,1.0)
    avail_w=RW-2*MARGIN_PX; avail_h=RH-2*MARGIN_PX
    sc=min(SCALE_BASE, avail_w/cW, avail_h/cH)
    CUR['sc']=sc
    CUR['ox']=(RW - cW*sc)/2 - minX*sc
    CUR['oy']=(RH - cH*sc)/2 - minY*sc

def ex(px): return Emu(int(round((CUR['ox'] + px*CUR['sc']) * EMU_PER_PX)))
def ey(px): return Emu(int(round((CUR['oy'] + px*CUR['sc']) * EMU_PER_PX)))
def ed(px): return Emu(int(round(max(px, 1) * CUR['sc'] * EMU_PER_PX)))

# 화면 순서/제목은 페이지리스트 네비 기준
def nav_order(html):
    out=[]
    for m in re.finditer(r'wf-pagelist-item[^>]*data-wf-action="(n\d+)"[^>]*>.*?wf-pagelist-name">([^<]*)<', html):
        out.append((m.group(1), m.group(2)))
    return out

EXTRACT_JS = r"""
(pid) => {
  const sec = document.querySelector(`section[data-page-id="${pid}"]`);
  const root = sec.querySelector('.wf-renderer');
  const rr = root.getBoundingClientRect();
  const toRGBA = (s) => {
    const m = s && s.match(/rgba?\(([^)]+)\)/);
    if(!m) return [0,0,0,0];
    const p = m[1].split(',').map(x=>parseFloat(x.trim()));
    return [p[0]||0, p[1]||0, p[2]||0, p.length>3?p[3]:1];
  };
  const out = [];
  let svgIdx = 0;
  const walk = (el, depth) => {
    const tag = el.tagName.toLowerCase();
    if (tag === 'svg') {
      const r = el.getBoundingClientRect();
      if (r.width>0 && r.height>0 && r.width<260) {
        el.setAttribute('data-wf-svg', svgIdx);
        out.push({kind:'svg', svg:svgIdx, depth,
                  x:r.left-rr.left, y:r.top-rr.top, w:r.width, h:r.height});
        svgIdx++;
      }
      return; // svg 내부는 더 안 들어감
    }
    const cs = getComputedStyle(el);
    const r = el.getBoundingClientRect();
    const x = r.left-rr.left, y = r.top-rr.top, w = r.width, h = r.height;
    const bg = toRGBA(cs.backgroundColor);
    const bw = parseFloat(cs.borderTopWidth)||0;
    const bc = toRGBA(cs.borderTopColor);
    const br = parseFloat(cs.borderTopLeftRadius)||0;
    const dashed = cs.borderTopStyle === 'dashed' || cs.borderTopStyle==='dotted';
    const isInput = (tag==='input' || tag==='textarea');
    // 직접 텍스트 노드만 수집(선행 아이콘 등은 자식 요소이므로 제외)
    let txt = ''; const tnodes=[];
    for (const n of el.childNodes) if (n.nodeType===3 && n.textContent.trim()){ txt += n.textContent; tnodes.push(n); }
    txt = txt.replace(/\s+/g,' ').trim();
    const col = toRGBA(cs.color);
    const fs = parseFloat(cs.fontSize)||14;
    const fw = parseInt(cs.fontWeight)||400;
    const align = cs.textAlign;
    let lh = parseFloat(cs.lineHeight); if (isNaN(lh)) lh = fs*1.4;
    if (w>0 && h>0) {
      const hasBox = (bg[3]>0.01) || (bw>0 && bc[3]>0.01) || isInput;
      if (hasBox) out.push({kind:'box', depth, x,y,w,h,
        bg, bw, bc, br, dashed, isInput});
      if (txt && tnodes.length) {
        // 텍스트 박스는 요소 전체가 아니라 실제 텍스트 범위(Range) 기준 → 선행 아이콘 뒤에서 시작
        const rg=document.createRange();
        rg.setStartBefore(tnodes[0]); rg.setEndAfter(tnodes[tnodes.length-1]);
        const tr=rg.getBoundingClientRect();
        const tx=tr.left-rr.left, ty=tr.top-rr.top, tw=tr.width, th=tr.height;
        const multiline=(lh>0 && th>lh*1.6);
        out.push({kind:'text', depth, x:tx,y:ty,w:tw,h:th, text:txt,
          col, fs, fw, align, multiline});
      }
      else if (isInput && el.placeholder) out.push({kind:'text', depth, x,y,w,h,
        text:el.placeholder, col:[150,150,150,1], fs, fw:400, align, multiline:false});
    }
    for (const c of el.children) walk(c, depth+1);
  };
  walk(root, 0);
  return out;
}
"""

def col(rgba):
    # 반투명색을 흰 슬라이드 배경 위로 블렌딩 (alpha 무시 방지)
    a = rgba[3] if len(rgba) > 3 else 1.0
    r = rgba[0]*a + 255*(1-a)
    g = rgba[1]*a + 255*(1-a)
    b = rgba[2]*a + 255*(1-a)
    return RGBColor(int(round(r)), int(round(g)), int(round(b)))

def set_dash(line, kind="dash"):
    ln = line._get_or_add_ln()
    d = ln.find(qn('a:prstDash'))
    if d is None:
        d = ln.makeelement(qn('a:prstDash'), {})
        ln.append(d)
    d.set('val', kind)

def add_box(slide, b):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if b['br']>0.5 else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, ex(b['x']), ey(b['y']), ed(b['w']), ed(b['h']))
    if shp_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = max(0.0, min(0.5, b['br']/max(min(b['w'],b['h']),1)))
        except Exception: pass
    f = s.fill
    if b['isInput']:
        f.solid(); f.fore_color.rgb = RGBColor(0xff,0xff,0xff)
    elif b['bg'][3] > 0.01:
        f.solid(); f.fore_color.rgb = col(b['bg'])
    else:
        f.background()
    ln = s.line
    if b['bw'] > 0 and b['bc'][3] > 0.01:
        ln.color.rgb = col(b['bc']); ln.width = Pt(max(b['bw']*0.75, 0.5))
        if b['dashed']: set_dash(ln, "dash")
    elif b['isInput']:
        ln.color.rgb = RGBColor(0xd0,0xd0,0xd0); ln.width = Pt(0.75)
    else:
        ln.fill.background()
    s.shadow.inherit = False
    return s

def add_text(slide, b):
    tb = slide.shapes.add_textbox(ex(b['x']), ey(b['y']), ed(b['w']), ed(b['h']))
    tf = tb.text_frame
    # HTML에서 여러 줄이던 텍스트만 줄바꿈(긴 문장은 박스 안에서 줄바꿈), 한 줄이던 라벨은 한 줄 유지
    tf.word_wrap = bool(b.get('multiline'))
    tf.auto_size = None
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ('margin_left','margin_right','margin_top','margin_bottom'):
        setattr(tf, side, Emu(0))
    p = tf.paragraphs[0]
    p.alignment = {'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT,
                   'justify':PP_ALIGN.JUSTIFY}.get(b['align'], PP_ALIGN.LEFT)
    run = p.add_run(); run.text = b['text']
    fnt = run.font
    fnt.size = Pt(max(b['fs']*FONT_PT_PER_PX*CUR['sc'], 5))
    fnt.bold = b['fw'] >= 600
    fnt.color.rgb = col(b['col'])
    fnt.name = KFONT
    return tb

def main():
    html = open(SRC, encoding="utf-8").read()
    order = nav_order(html)
    print(f"화면 {len(order)}개 변환 시작")

    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W_IN*914400))
    prs.slide_height = Emu(int(SLIDE_W_IN*914400 * RH/RW))
    blank = prs.slide_layouts[6]

    with sync_playwright() as pw:
        browser = pw.chromium.launch(channel="chrome", headless=True)
        page = browser.new_page(viewport={"width":1480,"height":940})
        page.goto("file://" + os.path.abspath(SRC))
        # 페이지리스트 네비 숨김
        page.evaluate("""() => { const n=document.querySelector('.wf-pagelist'); if(n) n.style.display='none'; }""")
        # 메뉴(대시보드/디바이스/알림/관리자) 항목 사이 세로 여백 확대 (사이드바는 빈 공간이 충분)
        page.add_style_tag(content="""
          .wf-renderer div[style*="width:240px"][style*="gap:8px"]{
            gap:22px !important; padding-top:26px !important;
          }
        """)

        for i,(pid,name) in enumerate(order, 1):
            # 이 섹션만 보이게
            page.evaluate("""(pid)=>{
                document.querySelectorAll('section[data-page-id]').forEach(s=>{
                    s.hidden = (s.getAttribute('data-page-id')!==pid);
                    s.style.display = (s.getAttribute('data-page-id')===pid)?'block':'none';
                });
                window.scrollTo(0,0);
            }""", pid)
            page.wait_for_timeout(60)
            boxes = page.evaluate(EXTRACT_JS, pid)

            slide = prs.slides.add_slide(blank)
            # 슬라이드 제목 = 노트
            slide.notes_slide.notes_text_frame.text = f"{i}. {name} ({pid})"

            fit_screen(boxes)   # 이 화면 콘텐츠를 슬라이드에 맞춤(경계 이탈 방지)
            order_key = {'box':0, 'svg':1, 'text':2}
            boxes.sort(key=lambda b:(b['depth'], order_key[b['kind']]))
            nsvg=0
            for b in boxes:
                if b['kind']=='box':
                    add_box(slide, b)
                elif b['kind']=='text':
                    add_text(slide, b)
                elif b['kind']=='svg':
                    # 셀렉터를 현재 섹션으로 한정(인덱스가 섹션마다 재시작하므로 중복 방지)
                    el = page.query_selector(f'section[data-page-id="{pid}"] [data-wf-svg="{b["svg"]}"]')
                    if el:
                        try:
                            png = el.screenshot(omit_background=True, timeout=3000)
                            slide.shapes.add_picture(io.BytesIO(png), ex(b['x']), ey(b['y']),
                                                     ed(b['w']), ed(b['h']))
                            nsvg+=1
                        except Exception: pass
            print(f"  [{i:2}/{len(order)}] {pid} {name:18} 도형 {len(boxes):3} (svg {nsvg})")
        browser.close()

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(f"\n저장: {OUT}  (슬라이드 {len(prs.slides._sldIdLst)})")

if __name__ == "__main__":
    main()
