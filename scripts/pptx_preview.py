#!/usr/bin/env python3
"""지정한 슬라이드들을 실제 pptx에서 읽어 HTML로 미리보기 렌더.
검증용 — 크롬에서 눈으로 확인."""
import sys, html, re
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn

PPTX = "발표자료/와이어프레임_v2.pptx"
SCALE = 72.0  # 1 inch -> px

def emu_in(v): return Emu(v).inches

def shape_fill(sh):
    el = sh._element
    sp = el.find(qn('p:spPr'))
    if sp is None: return None, None
    sf = sp.find(qn('a:solidFill'))
    if sf is None: return None, None
    srgb = sf.find(qn('a:srgbClr'))
    if srgb is None: return None, None
    col = srgb.get('val')
    alpha = srgb.find(qn('a:alpha'))
    a = int(alpha.get('val'))/100000 if alpha is not None else 1.0
    return col, a

def shape_line(sh):
    sp = sh._element.find(qn('p:spPr'))
    if sp is None: return None
    ln = sp.find(qn('a:ln'))
    if ln is None: return None
    sf = ln.find(qn('a:solidFill'))
    if sf is None: return None
    srgb = sf.find(qn('a:srgbClr'))
    return srgb.get('val') if srgb is not None else None

def shape_radius(sh):
    av = sh._element.findall('.//'+qn('a:gd'))
    for g in av:
        if g.get('name')=='adj':
            m=re.search(r'val (\d+)', g.get('fmla',''))
            if m: return int(m.group(1))/100000
    return 0

def runs(sh):
    out=[]
    if not sh.has_text_frame: return out
    for p in sh.text_frame.paragraphs:
        align = p.alignment
        for r in p.runs:
            sz = r.font.size.pt if r.font.size else 11
            col = r.font.color.rgb if (r.font.color and r.font.color.type is not None) else None
            out.append((r.text, sz, str(col) if col else "333333", bool(r.font.bold), str(align)))
    return out

def render(indices, outpath):
    pr = Presentation(PPTX)
    W = emu_in(pr.slide_width)*SCALE
    H = emu_in(pr.slide_height)*SCALE
    parts=["<html><head><meta charset='utf-8'><style>",
           "body{background:#444;font-family:'Apple SD Gothic Neo',sans-serif;margin:0;padding:24px}",
           ".cap{color:#fff;margin:18px 0 6px;font-size:14px}",
           f".slide{{position:relative;width:{W:.0f}px;height:{H:.0f}px;background:#fff;box-shadow:0 4px 18px rgba(0,0,0,.4);overflow:hidden}}",
           ".sh{position:absolute;box-sizing:border-box;display:flex;overflow:hidden}",
           "</style></head><body>"]
    for idx in indices:
        s = pr.slides[idx]
        parts.append(f"<div class='cap'>슬라이드 {idx+1}</div><div class='slide'>")
        for sh in s.shapes:
            try:
                l=emu_in(sh.left)*SCALE; t=emu_in(sh.top)*SCALE
                w=emu_in(sh.width)*SCALE; h=emu_in(sh.height)*SCALE
            except: continue
            col,a = shape_fill(sh)
            line = shape_line(sh)
            rad = shape_radius(sh)
            style=f"left:{l:.1f}px;top:{t:.1f}px;width:{w:.1f}px;height:{h:.1f}px;"
            if col:
                style+=f"background:rgba({int(col[0:2],16)},{int(col[2:4],16)},{int(col[4:6],16)},{a:.2f});"
            if line:
                style+=f"border:1px solid #{line};"
            if rad: style+=f"border-radius:{rad*100:.0f}%;" if rad>0.3 else f"border-radius:{min(rad*h*2,18):.0f}px;"
            rr = runs(sh)
            inner=""
            if rr:
                txt,sz,tc,bold,align = rr[0]
                full=" ".join(x[0] for x in rr)
                jc={'PP_ALIGN.CENTER':'center','PP_ALIGN.RIGHT':'flex-end'}.get(align,'flex-start')
                style+=f"align-items:center;justify-content:{jc};padding:0 3px;"
                inner=f"<span style='font-size:{sz:.1f}px;color:#{tc};font-weight:{700 if bold else 400};line-height:1.2'>{html.escape(full)}</span>"
            parts.append(f"<div class='sh' style='{style}'>{inner}</div>")
        parts.append("</div>")
    parts.append("</body></html>")
    with open(outpath,"w") as f: f.write("\n".join(parts))
    print("wrote", outpath)

if __name__=="__main__":
    idxs=[int(x) for x in sys.argv[1].split(",")] if len(sys.argv)>1 else [2,3,4]
    out=sys.argv[2] if len(sys.argv)>2 else "/tmp/wf_preview.html"
    render(idxs, out)
