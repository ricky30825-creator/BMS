#!/usr/bin/env python3
"""
와이어프레임_v2.pptx에 슬라이드 26(계정 설정) + 팝업 4장 추가.
  26: 계정 설정 기본 화면 (아이디·비밀번호·전화번호·이메일 각 수정하기 버튼)
  27: 아이디 수정 팝업
  28: 비밀번호 변경 팝업 (입력 3개)
  29: 전화번호 수정 팝업
  30: 이메일 수정 팝업
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

PPTX = "발표자료/와이어프레임_v2.pptx"
FONT = "Apple SD Gothic Neo"

C = dict(
    ink=RGBColor(0x33, 0x33, 0x33),
    sub=RGBColor(0x99, 0x99, 0x99),
    placeholder=RGBColor(0x96, 0x96, 0x96),
    white=RGBColor(0xFF, 0xFF, 0xFF),
    border=RGBColor(0xD0, 0xD0, 0xD0),
    cardborder=RGBColor(0xE6, 0xE6, 0xE6),
    btn_cancel_bg=RGBColor(0xF2, 0xF2, 0xF2),
    btn_cancel_bd=RGBColor(0xCB, 0xCA, 0xCA),
    dark=RGBColor(0x33, 0x33, 0x33),
    accent=RGBColor(0x6C, 0x5C, 0xE7),
)


def _set_run(r, text, size_pt, color, bold=False):
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.color.rgb = color


def add_text(sl, l, t, w, h, text, size_pt, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size_pt, color, bold)
    return tb


def add_rrect(sl, l, t, w, h, fill, line=None, line_w=0.75, radius=0.12):
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(l), Inches(t), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    return shp


def add_button(sl, l, t, w, h, text, primary=False):
    bg = C['dark'] if primary else C['btn_cancel_bg']
    bd = C['dark'] if primary else C['btn_cancel_bd']
    tx = C['white'] if primary else C['ink']
    shp = add_rrect(sl, l, t, w, h, bg, bd, 0.75, radius=0.16)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), text, 7.2, tx, bold=False)
    return shp


def add_input(sl, l, t, w, h, placeholder):
    add_rrect(sl, l, t, w, h, C['white'], C['border'], 0.75, radius=0.14)
    add_text(sl, l + 0.14, t, w - 0.24, h, placeholder, 7.0,
             C['placeholder'], anchor=MSO_ANCHOR.MIDDLE)


def copy_slide(pr, src):
    new = pr.slides.add_slide(src.slide_layout)
    for ph in list(new.shapes):
        ph._element.getparent().remove(ph._element)
    src_tree = src.shapes._spTree
    new_tree = new.shapes._spTree
    for el in src_tree:
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
            continue
        new_tree.append(copy.deepcopy(el))
    for blip in new_tree.findall('.//' + qn('a:blip')):
        embed = blip.get(qn('r:embed'))
        if embed:
            part = src.part.related_part(embed)
            blip.set(qn('r:embed'), new.part.relate_to(part, RT.IMAGE))
    return new


def _in(v):
    return Emu(v).inches if v else 0


def build_account_settings(sl):
    """
    슬라이드에서 콘텐츠 영역(x>2.9" 이고 y>0.86") 셰이프만 제거하고
    계정 설정 내용을 새로 그린다.
    """
    # 콘텐츠 영역 셰이프 제거 (헤더·사이드바는 보존)
    to_remove = []
    for sh in sl.shapes:
        try:
            l, t = _in(sh.left), _in(sh.top)
        except Exception:
            continue
        if t > 0.86 and l > 2.9:
            to_remove.append(sh)
    for sh in to_remove:
        sh._element.getparent().remove(sh._element)

    # ── 상단 서브 탭 (릴레이 제어 | 알림 설정 | 계정 설정) ──────────
    add_text(sl, 3.05, 1.05, 0.65, 0.16, "릴레이 제어", 7.5, C['sub'])
    add_text(sl, 3.84, 1.05, 0.65, 0.16, "알림 설정", 7.5, C['sub'])
    add_text(sl, 4.63, 1.04, 0.58, 0.16, "계정 설정",   7.5, C['ink'], bold=True)
    # 활성 탭 밑줄
    add_rrect(sl, 4.63, 1.18, 0.58, 0.02, C['dark'], None, radius=0.0)

    # ── 계정 정보 4개 행 ──────────────────────────────────────────────
    ROW_H   = 0.82
    ROW_GAP = 0.12
    Y0      = 1.33
    CX      = 3.05   # 콘텐츠 영역 왼쪽
    CW      = 9.45   # 콘텐츠 영역 너비

    rows = [
        ("아이디",    "user_example"),
        ("비밀번호",  "••••••••••"),
        ("전화번호",  "010-****-****"),
        ("이메일",    "user@example.com"),
    ]

    for i, (label, value) in enumerate(rows):
        y = Y0 + i * (ROW_H + ROW_GAP)
        # 행 카드
        add_rrect(sl, CX, y, CW, ROW_H, C['white'], C['cardborder'], 0.75, radius=0.05)
        # 레이블 (항목명)
        add_text(sl, CX + 0.22, y + 0.12, 1.80, 0.16, label, 7.5, C['sub'])
        # 현재 값
        add_text(sl, CX + 0.22, y + 0.42, 7.00, 0.22, value, 9.4, C['ink'])
        # 수정하기 버튼 (오른쪽 정렬)
        btn_x = CX + CW - 0.22 - 0.88
        btn_y = y + (ROW_H - 0.30) / 2
        add_button(sl, btn_x, btn_y, 0.88, 0.30, "수정하기", primary=False)

    # 페이지 설명 안내
    add_text(sl, 3.05, 5.20, 9.45, 0.18,
             "개인정보 수정은 각 항목의 수정하기 버튼을 눌러 팝업에서 진행합니다.",
             6.2, C['sub'])


def build_edit_popup(sl, title, fields, confirm_label):
    """
    fields: [(label, placeholder), ...]
    confirm_label: 확인 버튼 텍스트
    """
    # 오버레이 (검정 45%)
    ov = add_rrect(sl, 0.63, 0.33, 12.07, 7.67, RGBColor(0, 0, 0), None, radius=0.0)
    srgb = ov.fill.fore_color._xFill.find(qn('a:srgbClr'))
    srgb.append(srgb.makeelement(qn('a:alpha'), {'val': '45000'}))

    n   = len(fields)
    MH  = 1.85 + n * 0.70        # 입력 개수에 따라 높이 조정
    ML  = 4.17
    MT  = max(1.80, 4.00 - MH / 2)   # 세로 중앙 정렬
    MW  = 5.00
    pad = 0.29
    ix  = ML + pad
    iw  = MW - 2 * pad

    # 모달 카드
    add_rrect(sl, ML, MT, MW, MH, C['white'], C['border'], 0.75, radius=0.05)

    # 제목
    add_text(sl, ix, MT + 0.25, iw, 0.21, title, 9.4, C['ink'], bold=True)

    # 입력 필드들
    for j, (flabel, placeholder) in enumerate(fields):
        fy = MT + 0.62 + j * 0.72
        add_text(sl, ix, fy, iw, 0.16, flabel, 7.0, C['ink'])
        add_input(sl, ix, fy + 0.20, iw, 0.33, placeholder)

    # 버튼 (우하단)
    by       = MT + MH - 0.33 - 0.20
    bx_right = ML + MW - pad
    cw       = 0.82   # 확인 버튼 너비
    add_button(sl, bx_right - cw,              by, cw,   0.33, confirm_label, primary=True)
    add_button(sl, bx_right - cw - 0.10 - 0.57, by, 0.57, 0.33, "취소",       primary=False)


def main():
    pr = Presentation(PPTX)

    # 슬라이드 25를 베이스로 계정 설정 슬라이드 생성
    base = pr.slides[24]
    sl26 = copy_slide(pr, base)
    build_account_settings(sl26)

    # 팝업 4장 정의
    popups = [
        ("아이디 수정",
         [("새 아이디", "새 아이디를 입력하세요")],
         "수정 완료"),
        ("비밀번호 변경",
         [("현재 비밀번호", "현재 비밀번호 입력"),
          ("새 비밀번호",   "새 비밀번호 입력"),
          ("비밀번호 확인", "새 비밀번호 재입력")],
         "변경 완료"),
        ("전화번호 수정",
         [("새 전화번호", "010-0000-0000")],
         "수정 완료"),
        ("이메일 수정",
         [("새 이메일 주소", "예: user@example.com")],
         "수정 완료"),
    ]

    popup_slides = []
    for title, fields, confirm in popups:
        sl = copy_slide(pr, sl26)
        build_edit_popup(sl, title, fields, confirm)
        popup_slides.append(sl)

    pr.save(PPTX)
    print(f"saved. total slides: {len(pr.slides._sldIdLst)}")


if __name__ == "__main__":
    main()
