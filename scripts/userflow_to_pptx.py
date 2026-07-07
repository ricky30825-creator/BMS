#!/usr/bin/env python3
"""슬라이드 11(Flow chart)의 기존 이미지를 제거하고
업데이트된 유저플로우를 '수정 가능한 네이티브 도형'으로 생성한다.

스타일: 보라 스윔레인(HTML 버전과 동일)
 - start  : 검정      / 흰 글자
 - main   : 보라 단색  / 흰 글자
 - accent : 연보라     / 보라 글자 (게이트/서브화면/위험/진입)
 - action : 흰색       / 회색 테두리
엣지: 엘보 커넥터(편집 가능) + 화살표
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PPTX = os.path.join(ROOT, '발표자료', '설계 산출물_정재환 2.pptx')

# ── 다이어그램 데이터 (HTML 버전과 동일) ───────────────────────────
NW, NH = 176, 64          # 박스 확대(글자 넘침 방지)
LANES = ['인증', '디바이스 연동', '실시간 대시보드', 'AI 이상탐지', '알림 센터', '관리자 콘솔']

N = {
 'start':{'cx':165,'cy':110,'lane':1,'s':'start','w':78,'l':'시작'},
 'login':{'cx':340,'cy':110,'lane':1,'s':'main','l':'로그인 화면'},
 'loginNorm':{'cx':520,'cy':70,'lane':1,'s':'action','l':'일반 로그인'},
 'loginSSO':{'cx':520,'cy':150,'lane':1,'s':'action','l':'소셜 로그인(SSO)'},
 'role':{'cx':700,'cy':110,'lane':1,'s':'accent','l':'역할 확인(gate)'},
 'findAcc':{'cx':520,'cy':210,'lane':1,'s':'accent','l':'계정 찾기'},
 'findId':{'cx':705,'cy':188,'lane':1,'s':'action','l':'아이디 찾기'},
 'resetPw':{'cx':705,'cy':232,'lane':1,'s':'action','l':'비밀번호 재설정'},

 'device':{'cx':165,'cy':362,'lane':2,'s':'main','l':'디바이스 연동'},
 'batSel':{'cx':355,'cy':362,'lane':2,'s':'accent','l':'배터리 선택/등록'},
 'batSaved':{'cx':548,'cy':305,'lane':2,'s':'action','l':'저장된 배터리 선택'},
 'batNew':{'cx':548,'cy':362,'lane':2,'s':'action','l':'새 배터리 등록'},
 'batHist':{'cx':548,'cy':419,'lane':2,'s':'accent','l':'배터리 상세/이력'},
 'interlock':{'cx':742,'cy':362,'lane':2,'s':'accent','l':'모드 인터락 차단(gate)'},
 'mode1':{'cx':930,'cy':305,'lane':2,'s':'action','l':'모드1 내장배터리'},
 'mode2':{'cx':930,'cy':362,'lane':2,'s':'action','l':'모드2 외부 셀'},
 'mode3':{'cx':930,'cy':419,'lane':2,'s':'action','l':'모드3 보조배터리'},
 'measure':{'cx':1120,'cy':362,'lane':2,'s':'action','l':'측정 시작(세션 생성)'},
 'sensor':{'cx':1305,'cy':362,'lane':2,'s':'action','l':'센서 수집 상태'},
 'collectErr':{'cx':1305,'cy':300,'lane':2,'s':'accent','l':'수집 오류 이벤트'},

 'dash':{'cx':165,'cy':650,'lane':3,'s':'main','l':'실시간 대시보드'},
 'logout':{'cx':380,'cy':528,'lane':3,'s':'action','l':'로그아웃'},
 'gauge':{'cx':380,'cy':582,'lane':3,'s':'action','l':'V/I/T/SOC 게이지'},
 'anomShow':{'cx':380,'cy':636,'lane':3,'s':'action','l':'이상점수 표시'},
 'warnBanner':{'cx':380,'cy':690,'lane':3,'s':'accent','l':'경고 알림 배너'},
 'trend':{'cx':380,'cy':744,'lane':3,'s':'action','l':'추세 차트/이력 조회'},
 'adminGate':{'cx':380,'cy':798,'lane':3,'s':'accent','l':'관리자 권한 확인(gate)'},
 'period':{'cx':578,'cy':744,'lane':3,'s':'action','l':'기간 조회'},

 'ai':{'cx':1120,'cy':905,'lane':4,'s':'main','l':'AI 이상탐지'},
 'score':{'cx':1305,'cy':905,'lane':4,'s':'action','l':'이상점수 산출'},
 'judge':{'cx':1490,'cy':905,'lane':4,'s':'action','l':'정상/주의/경고 판정'},
 'danger':{'cx':1690,'cy':852,'lane':4,'s':'accent','l':'위험 등급 감지(danger)'},
 'evStore':{'cx':1690,'cy':958,'lane':4,'s':'action','l':'이상 이벤트 저장'},
 'killAuto':{'cx':1885,'cy':856,'lane':4,'s':'accent','l':'자동 Kill-Switch 차단'},
 'alertNow':{'cx':1885,'cy':954,'lane':4,'s':'action','l':'즉시 알림 발송'},

 'katalk':{'cx':165,'cy':1095,'lane':5,'s':'entry','l':'카톡 알림 진입(entry)'},
 'notif':{'cx':1490,'cy':1095,'lane':5,'s':'main','l':'알림 센터'},
 'sns':{'cx':1690,'cy':1095,'lane':5,'s':'action','l':'SNS 알림 발송'},
 'notifHist':{'cx':1885,'cy':1095,'lane':5,'s':'action','l':'알림 이력 조회'},
 'notifDet':{'cx':2080,'cy':1095,'lane':5,'s':'action','l':'알림 상세 확인'},

 'admin':{'cx':742,'cy':1245,'lane':6,'s':'main','l':'관리자 콘솔'},
 'permMgmt':{'cx':930,'cy':1195,'lane':6,'s':'action','l':'권한 관리'},
 'relayRec':{'cx':930,'cy':1295,'lane':6,'s':'action','l':'릴레이 복구'},
 'permGate':{'cx':1120,'cy':1195,'lane':6,'s':'accent','l':'권한변경 확인(gate)'},
 'permApply':{'cx':1305,'cy':1195,'lane':6,'s':'action','l':'권한 변경 적용'},
 'recGate':{'cx':1120,'cy':1295,'lane':6,'s':'accent','l':'복구 확인(gate)'},
 'relayBlock':{'cx':1490,'cy':1245,'lane':6,'s':'action','l':'릴레이 차단'},
 'blockGate':{'cx':1680,'cy':1245,'lane':6,'s':'accent','l':'차단 확인(gate)'},
}

E = [
 ('start','login'),('login','loginNorm'),('login','loginSSO'),('login','findAcc'),
 ('loginNorm','role'),('loginSSO','role'),('findAcc','findId'),('findAcc','resetPw'),
 ('role','device'),
 ('device','batSel'),('batSel','batSaved'),('batSel','batNew'),('batSel','batHist'),
 ('batSaved','interlock'),('interlock','mode1'),('interlock','mode2'),('interlock','mode3'),
 ('mode1','measure'),('mode2','measure'),('mode3','measure'),
 ('measure','sensor'),('sensor','collectErr'),('measure','dash'),
 ('dash','logout'),('dash','gauge'),('dash','anomShow'),('dash','warnBanner'),
 ('dash','trend'),('dash','adminGate'),('trend','period'),
 ('dash','ai'),('dash','notif'),('adminGate','admin'),
 ('ai','score'),('score','judge'),('judge','danger'),('judge','evStore'),
 ('danger','killAuto'),('danger','alertNow'),('alertNow','notif'),
 ('notif','sns'),('sns','notifHist'),('notifHist','notifDet'),
 ('katalk','dash','deeplink'),
 ('admin','permMgmt'),('admin','relayRec'),
 ('permMgmt','permGate'),('permGate','permApply'),
 ('relayRec','recGate'),('admin','relayBlock'),('relayBlock','blockGate'),
]

# 간격 확대 스케일(박스 확대에 맞춰 간격도 키움)
SX, SY = 1.34, 1.50
for k in N:
    N[k]['cx'] = N[k]['cx'] * SX
    N[k]['cy'] = N[k]['cy'] * SY

def w(n): return n.get('w', NW)

# ── 색상 ───────────────────────────────────────────────────────────
C = {
 'start_bg':RGBColor(0x2B,0x2B,0x33),'start_tx':RGBColor(0xFF,0xFF,0xFF),
 'main_bg':RGBColor(0x6C,0x5C,0xE7),'main_tx':RGBColor(0xFF,0xFF,0xFF),
 'accent_bg':RGBColor(0xEC,0xE8,0xFF),'accent_bd':RGBColor(0xC7,0xBC,0xF5),'accent_tx':RGBColor(0x5A,0x49,0xCE),
 'action_bg':RGBColor(0xFF,0xFF,0xFF),'action_bd':RGBColor(0xD7,0xD7,0xE0),'action_tx':RGBColor(0x4A,0x4A,0x57),
 'edge':RGBColor(0xB6,0xB6,0xC4),'lane_tx':RGBColor(0x9A,0x9A,0xAC),
 'band':RGBColor(0x6C,0x5C,0xE7),'lane_line':RGBColor(0xEC,0xEC,0xF2),
}

# ── 배치 영역(슬라이드 EMU) ────────────────────────────────────────
BOX_L, BOX_T = Inches(1.0), Inches(1.18)
BOX_W, BOX_H = Inches(8.9), Inches(5.92)   # 노드 영역(레인 라벨은 왼쪽 여백 사용)

xs0 = min(n['cx'] - w(n)/2 for n in N.values())
xs1 = max(n['cx'] + w(n)/2 for n in N.values())
ys0 = min(n['cy'] - NH/2 for n in N.values())
ys1 = max(n['cy'] + NH/2 for n in N.values())
diagW, diagH = xs1 - xs0, ys1 - ys0
scale = min(BOX_W / diagW, BOX_H / diagH)          # EMU per px
offX = BOX_L + int((BOX_W - diagW * scale) / 2)
offY = BOX_T + int((BOX_H - diagH * scale) / 2)

def mx(px): return int(offX + (px - xs0) * scale)
def my(px): return int(offY + (px - ys0) * scale)

# ── 엣지 경로 끝점(HTML border 로직) ───────────────────────────────
def endpoints(a, b):
    aL, aR = a['cx']-w(a)/2, a['cx']+w(a)/2
    aT, aB = a['cy']-NH/2, a['cy']+NH/2
    bL, bR = b['cx']-w(b)/2, b['cx']+w(b)/2
    bT, bB = b['cy']-NH/2, b['cy']+NH/2
    if abs(b['cy']-a['cy']) <= 70:        # 같은 레인: 수평
        right = b['cx'] >= a['cx']
        return ((aR if right else aL), a['cy'], (bL if right else bR), b['cy'])
    else:                                 # 레인 간: 수직
        down = b['cy'] >= a['cy']
        return (a['cx'], (aB if down else aT), b['cx'], (bT if down else bB))

def set_line(ln_parent_line, color, width_pt, arrow=True, dash=False):
    """LineFormat 설정 + 화살표/점선(XML)."""
    ln_parent_line.color.rgb = color
    ln_parent_line.width = Pt(width_pt)
    ln = ln_parent_line._get_or_add_ln()
    if dash:
        pd = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(pd)
    if arrow:
        te = ln.makeelement(qn('a:tailEnd'), {'type':'triangle','w':'med','len':'med'})
        ln.append(te)


def build(slide):
    shapes = slide.shapes

    # 1) 기존 이미지 제거
    for sh in list(shapes):
        if sh.shape_type == 13:  # PICTURE
            sh._element.getparent().remove(sh._element)

    # 2) 레인 밴드 + 라벨 (뒤)
    lane_y = {}
    for n in N.values():
        b = lane_y.setdefault(n['lane'], [1e9, -1e9])
        b[0] = min(b[0], n['cy']-NH/2); b[1] = max(b[1], n['cy']+NH/2)
    order = sorted(lane_y)
    for i, ln in enumerate(order):
        cur = lane_y[ln]
        top = ys0 if i == 0 else (lane_y[order[i-1]][1] + cur[0]) / 2
        bot = ys1 if i == len(order)-1 else (cur[1] + lane_y[order[i+1]][0]) / 2
        if i % 2 == 1:
            band = shapes.add_shape(MSO_SHAPE.RECTANGLE, mx(xs0), my(top),
                                    int(diagW*scale), my(bot)-my(top))
            band.fill.solid(); band.fill.fore_color.rgb = RGBColor(0xF4,0xF2,0xFD)
            band.line.fill.background()
            band.shadow.inherit = False
        # 라벨
        lbl = shapes.add_textbox(Inches(0.28), my((cur[0]+cur[1])/2)-Inches(0.12),
                                 Inches(0.95), Inches(0.3))
        tf = lbl.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; r = p.add_run(); r.text = LANES[ln-1]
        r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = C['lane_tx']

    # 3) 엣지 (가운데)
    for e in E:
        a, b = N[e[0]], N[e[1]]
        dash = len(e) > 2 and e[2] == 'deeplink'
        x1, y1, x2, y2 = endpoints(a, b)
        cxn = shapes.add_connector(MSO_CONNECTOR.ELBOW, mx(x1), my(y1), mx(x2), my(y2))
        set_line(cxn.line, C['edge'], 0.75, arrow=True, dash=dash)

    # 4) 노드 (앞)
    for n in N.values():
        st = n['s']
        shp = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               mx(n['cx']-w(n)/2), my(n['cy']-NH/2),
                               int(w(n)*scale), int(NH*scale))
        shp.adjustments[0] = 0.22
        shp.shadow.inherit = False
        fill = shp.fill; fill.solid()
        line = shp.line
        if st == 'start':
            fill.fore_color.rgb = C['start_bg']; line.fill.background(); tx = C['start_tx']; bold = True
        elif st == 'main':
            fill.fore_color.rgb = C['main_bg']; line.fill.background(); tx = C['main_tx']; bold = True
        elif st in ('accent', 'entry'):
            fill.fore_color.rgb = C['accent_bg']; line.color.rgb = C['accent_bd']; line.width = Pt(0.75)
            tx = C['accent_tx']; bold = False
            if st == 'entry':
                ln = line._get_or_add_ln()
                ln.append(ln.makeelement(qn('a:prstDash'), {'val':'dash'}))
        else:  # action
            fill.fore_color.rgb = C['action_bg']; line.color.rgb = C['action_bd']; line.width = Pt(0.75)
            tx = C['action_tx']; bold = False
        tf = shp.text_frame; tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE   # 넘치면 글자 자동 축소
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = n['l']
        r.font.size = Pt(7); r.font.bold = bold; r.font.color.rgb = tx


def main():
    pr = Presentation(PPTX)
    slide = pr.slides[10]  # 11페이지
    build(slide)
    pr.save(PPTX)
    print('saved:', PPTX)
    print(f'nodes={len(N)} edges={len(E)} scale={scale:.6f} '
          f'area=({BOX_W/914400:.2f}x{BOX_H/914400:.2f}in)')

if __name__ == '__main__':
    main()
