#!/usr/bin/env python3
"""와이어프레임_v2.pptx 수정:
 1) 슬라이드 3 '계정 찾기'를 '아이디 찾기 / 비밀번호 재설정' 2개 카드 선택 화면으로 재구성
    (배터리 선택/등록 화면과 동일한 패턴, 각 카드에 '▸ 팝업으로 열림' 표기)
 2) 슬라이드 3 바로 뒤에 팝업 슬라이드 2장 추가
    - '아이디 찾기' 팝업 (어두운 오버레이 + 중앙 모달, 배터리 팝업과 동일 스타일)
    - '비밀번호 재설정' 팝업
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

PPTX = "발표자료/와이어프레임_v2.pptx"
FONT = "Apple SD Gothic Neo"

# ── 색상 ─────────────────────────────────────────────────────
C = dict(
    ink=RGBColor(0x33, 0x33, 0x33),
    sub=RGBColor(0x99, 0x99, 0x99),
    placeholder=RGBColor(0x96, 0x96, 0x96),
    white=RGBColor(0xFF, 0xFF, 0xFF),
    border=RGBColor(0xD0, 0xD0, 0xD0),
    iconbg=RGBColor(0xF2, 0xF2, 0xF2),
    cardborder=RGBColor(0xE6, 0xE6, 0xE6),
    btn_cancel_bg=RGBColor(0xF2, 0xF2, 0xF2),
    btn_cancel_bd=RGBColor(0xCB, 0xCA, 0xCA),
    dark=RGBColor(0x33, 0x33, 0x33),
    accent=RGBColor(0x6C, 0x5C, 0xE7),
)


def _set_run(r, text, size_pt, color, bold=False):
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.color.rgb = color


def add_text(sl, l, t, w, h, text, size_pt, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size_pt, color, bold)
    return tb


def add_rrect(sl, l, t, w, h, fill, line=None, line_w=0.75, radius=0.12):
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(l), Inches(t), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    return shp


def add_button(sl, l, t, w, h, text, primary=False):
    bg = C['dark'] if primary else C['btn_cancel_bg']
    bd = C['dark'] if primary else C['btn_cancel_bd']
    tx = C['white'] if primary else C['ink']
    shp = add_rrect(sl, l, t, w, h, bg, bd, 0.75, radius=0.16)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), text, 7.2, tx, bold=False)
    return shp


def add_input(sl, l, t, w, h, placeholder):
    add_rrect(sl, l, t, w, h, C['white'], C['border'], 0.75, radius=0.14)
    tb = add_text(sl, l + 0.14, t, w - 0.24, h, placeholder, 7.0,
                  C['placeholder'], anchor=MSO_ANCHOR.MIDDLE)
    return tb


# ── 슬라이드 복제 ─────────────────────────────────────────────
def copy_slide(pr, src):
    new = pr.slides.add_slide(src.slide_layout)
    for ph in list(new.shapes):
        ph._element.getparent().remove(ph._element)
    src_tree = src.shapes._spTree
    new_tree = new.shapes._spTree
    for el in src_tree:
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
            continue
        new_tree.append(copy.deepcopy(el))
    # 이미지 관계(돋보기 아이콘 등) 재연결
    for blip in new_tree.findall('.//' + qn('a:blip')):
        embed = blip.get(qn('r:embed'))
        if embed:
            part = src.part.related_part(embed)
            blip.set(qn('r:embed'), new.part.relate_to(part, RT.IMAGE))
    return new


def build_base_chooser(sl):
    """슬라이드 3의 본문을 2-카드 선택 화면으로 교체 (크롬 #0-13 유지)."""
    shapes = list(sl.shapes)
    for sh in shapes[14:]:
        sh._element.getparent().remove(sh._element)

    # 제목 + 설명
    add_text(sl, 3.05, 1.05, 4.0, 0.20, "계정 찾기", 9.4, C['ink'], bold=True)
    add_text(sl, 3.05, 1.29, 9.45, 0.24,
             "가입한 계정 정보를 찾거나 비밀번호를 재설정하세요. 각 기능은 팝업으로 열립니다.",
             6.4, C['sub'])

    cards = [
        (3.05, "🆔", "아이디 찾기",
         "등록된 이메일로 가입된 아이디(사용자명)를 전송해드립니다."),
        (7.91, "🔑", "비밀번호 재설정",
         "등록된 이메일로 비밀번호 재설정 링크를 보내드립니다."),
    ]
    for cx, icon, title, desc in cards:
        add_rrect(sl, cx, 1.66, 4.60, 1.56, C['white'], C['cardborder'], 0.75, radius=0.05)
        add_rrect(sl, cx + 0.22, 1.87, 0.51, 0.51, C['iconbg'], None, radius=0.18)
        add_text(sl, cx + 0.22, 2.06, 0.51, 0.22, icon, 12, C['ink'],
                 align=PP_ALIGN.CENTER)
        add_text(sl, cx + 0.22, 2.49, 4.16, 0.20, title, 9.4, C['ink'], bold=True)
        add_text(sl, cx + 0.22, 2.73, 4.16, 0.33, desc, 6.4, C['sub'])
        add_text(sl, cx + 0.22, 3.01, 4.16, 0.16, "▸ 팝업으로 열림", 6.4, C['accent'])

    add_text(sl, 3.05, 3.45, 2.0, 0.16, "← 로그인으로 돌아가기", 7.0, C['sub'])


def build_popup(sl, title, desc, info, primary_label, primary_w):
    """기존 슬라이드(base) 위에 어두운 오버레이 + 중앙 모달 추가."""
    # 오버레이 (검정 45%)
    ov = add_rrect(sl, 0.63, 0.33, 12.07, 7.67, RGBColor(0, 0, 0), None, radius=0.0)
    ov_fill = ov.fill.fore_color._xFill  # solidFill
    srgb = ov_fill.find(qn('a:srgbClr'))
    a = srgb.makeelement(qn('a:alpha'), {'val': '45000'})
    srgb.append(a)

    # 모달 카드
    ML, MT, MW, MH = 4.17, 2.30, 5.00, 2.75
    add_rrect(sl, ML, MT, MW, MH, C['white'], C['border'], 0.75, radius=0.05)

    pad = 0.29
    ix = ML + pad
    iw = MW - 2 * pad
    add_text(sl, ix, MT + 0.25, iw, 0.21, title, 9.4, C['ink'], bold=True)
    add_text(sl, ix, MT + 0.52, iw, 0.26, desc, 6.4, C['sub'])
    add_text(sl, ix, MT + 0.92, iw, 0.16, "이메일", 7.0, C['ink'])
    add_input(sl, ix, MT + 1.12, iw, 0.33, "예: user@example.com")
    add_text(sl, ix, MT + 1.56, iw, 0.20, info, 6.0, C['sub'])

    # 버튼 (우하단)
    by = MT + MH - 0.33 - 0.20
    bx_right = ML + MW - pad
    add_button(sl, bx_right - primary_w, by, primary_w, 0.33, primary_label, primary=True)
    add_button(sl, bx_right - primary_w - 0.10 - 0.57, by, 0.57, 0.33, "취소", primary=False)


def main():
    pr = Presentation(PPTX)

    base = pr.slides[2]            # 슬라이드 3 (계정 찾기)
    build_base_chooser(base)

    p_id = copy_slide(pr, base)
    build_popup(p_id,
                "아이디 찾기",
                "등록된 이메일로 가입된 아이디(사용자명)를 전송해드립니다.",
                "ⓘ 가입 시 등록한 이메일 주소를 입력하세요.",
                "전송", 0.57)

    p_pw = copy_slide(pr, base)
    build_popup(p_pw,
                "비밀번호 재설정",
                "등록된 이메일로 비밀번호 재설정 링크를 보내드립니다.",
                "ⓘ 메일의 링크에서 새 비밀번호를 설정할 수 있습니다.",
                "재설정 링크 전송", 1.30)

    # 새 슬라이드를 슬라이드 3 바로 뒤(인덱스 3,4)로 이동
    lst = pr.slides._sldIdLst
    ids = list(lst)
    new_id, new_pw = ids[-2], ids[-1]
    lst.remove(new_id)
    lst.remove(new_pw)
    lst.insert(3, new_id)
    lst.insert(4, new_pw)

    pr.save(PPTX)
    print("saved. total slides:", len(pr.slides._sldIdLst))


if __name__ == "__main__":
    main()
